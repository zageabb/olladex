from __future__ import annotations

import csv
import re
import shutil
import subprocess
import tempfile
from copy import copy
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formula import Tokenizer
from openpyxl.formula.translate import Translator
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.table import Table, TableStyleInfo


_TABLE_NAME = re.compile(r"^[A-Za-z_][A-Za-z0-9_.]*$")


def mutate_xlsx_plus(path: Path, operations: list[dict[str, Any]]) -> None:
    """Unified XLSX/XLSM editor combining Olladex and Context Studio operations."""
    book = load_workbook(path, data_only=False, keep_vba=path.suffix.lower() == ".xlsm")
    try:
        for operation in operations:
            action = str(operation.get("action") or "")
            if action == "set_cell":
                sheet = _sheet(book, str(operation.get("sheet") or book.active.title))
                sheet[_cell(operation)] = _coerce(operation.get("value"))
            elif action == "set_cells":
                sheet = _sheet(book, str(operation.get("sheet") or book.active.title))
                values = operation.get("values") or {}
                if not isinstance(values, dict):
                    raise ValueError("set_cells requires a values object keyed by cell address")
                for coordinate, value in values.items():
                    sheet[str(coordinate).upper()] = _coerce(value)
            elif action == "set_cell_format":
                sheet = _sheet(book, str(operation.get("sheet") or book.active.title))
                cell = sheet[_cell(operation)]
                font_color = _hex(operation.get("font_color")) if "font_color" in operation else _font_rgb(cell)
                cell.font = Font(
                    name=str(operation.get("font_name") or cell.font.name or "Calibri"),
                    size=float(operation.get("font_size") or cell.font.sz or 11),
                    bold=bool(operation.get("bold")) if "bold" in operation else bool(cell.font.bold),
                    italic=bool(operation.get("italic")) if "italic" in operation else bool(cell.font.italic),
                    underline="single" if operation.get("underline") else (cell.font.underline if "underline" not in operation else None),
                    color=font_color or None,
                )
                if "fill_color" in operation:
                    fill = _hex(operation.get("fill_color"))
                    cell.fill = PatternFill(fill_type="solid", fgColor=fill) if fill else PatternFill()
                cell.alignment = Alignment(
                    horizontal=str(operation.get("horizontal") or cell.alignment.horizontal or "general"),
                    vertical=str(operation.get("vertical") or cell.alignment.vertical or "bottom"),
                    wrap_text=bool(operation.get("wrap_text")) if "wrap_text" in operation else bool(cell.alignment.wrap_text),
                )
                if "number_format" in operation:
                    cell.number_format = str(operation.get("number_format") or "General")
            elif action == "set_range_values":
                sheet = _sheet(book, str(operation.get("sheet") or book.active.title))
                anchor = sheet[_cell(operation, "start_cell")]
                values = operation.get("values")
                if not isinstance(values, list):
                    raise ValueError("set_range_values requires values")
                for row_offset, row in enumerate(values):
                    if not isinstance(row, list):
                        continue
                    for column_offset, value in enumerate(row):
                        sheet.cell(row=anchor.row + row_offset, column=anchor.column + column_offset, value=_coerce(value))
            elif action == "fill":
                sheet = _sheet(book, str(operation.get("sheet") or book.active.title))
                source = str(operation.get("source") or "").upper()
                target = str(operation.get("range") or "").upper()
                if not source or not target:
                    raise ValueError("fill requires source and range")
                source_cell = sheet[source]
                for row in sheet[target]:
                    for cell in row:
                        if isinstance(source_cell.value, str) and source_cell.value.startswith("="):
                            cell.value = Translator(source_cell.value, origin=source_cell.coordinate).translate_formula(cell.coordinate)
                        else:
                            cell.value = source_cell.value
                        cell._style = copy(source_cell._style)
                        cell.number_format = source_cell.number_format
            elif action == "format":
                sheet = _sheet(book, str(operation.get("sheet") or book.active.title))
                target = str(operation.get("range") or "").upper()
                style = operation.get("style") or {}
                if not target or not isinstance(style, dict):
                    raise ValueError("format requires range and style")
                side = Side(style="thin", color=str(style.get("border_color") or "D9E2F0").replace("#", ""))
                for row in sheet[target]:
                    for cell in row:
                        if "bold" in style or "italic" in style or style.get("font_color"):
                            cell.font = Font(
                                name=cell.font.name,
                                size=cell.font.size,
                                bold=bool(style.get("bold")) if "bold" in style else bool(cell.font.bold),
                                italic=bool(style.get("italic")) if "italic" in style else bool(cell.font.italic),
                                color=str(style.get("font_color") or _font_rgb(cell) or "000000").replace("#", ""),
                            )
                        if style.get("fill"):
                            cell.fill = PatternFill("solid", fgColor=str(style["fill"]).replace("#", ""))
                        if style.get("number_format"):
                            cell.number_format = str(style["number_format"])
                        if style.get("wrap") is not None or style.get("horizontal"):
                            cell.alignment = Alignment(
                                horizontal=str(style.get("horizontal") or cell.alignment.horizontal or "general"),
                                vertical=cell.alignment.vertical,
                                wrap_text=bool(style.get("wrap")) if style.get("wrap") is not None else bool(cell.alignment.wrap_text),
                            )
                        if style.get("border"):
                            cell.border = Border(left=side, right=side, top=side, bottom=side)
            elif action == "add_sheet":
                name = _sheet_name(operation.get("name"))
                if name in book.sheetnames:
                    raise ValueError(f"Worksheet already exists: {name}")
                book.create_sheet(title=name)
            elif action == "rename_sheet":
                sheet = _sheet(book, str(operation.get("sheet") or book.active.title))
                name = _sheet_name(operation.get("name"))
                if name in book.sheetnames and name != sheet.title:
                    raise ValueError(f"Worksheet already exists: {name}")
                sheet.title = name
            elif action == "delete_sheet":
                if len(book.worksheets) <= 1:
                    raise ValueError("Workbook must contain at least one worksheet")
                book.remove(_sheet(book, str(operation.get("sheet") or book.active.title)))
            elif action in {"insert_rows", "delete_rows"}:
                sheet = _sheet(book, str(operation.get("sheet") or book.active.title))
                index = max(1, int(operation.get("index", 1)))
                amount = max(1, int(operation.get("amount", 1)))
                sheet.insert_rows(index, amount) if action == "insert_rows" else sheet.delete_rows(index, amount)
            elif action in {"insert_columns", "delete_columns"}:
                sheet = _sheet(book, str(operation.get("sheet") or book.active.title))
                index = max(1, int(operation.get("index", 1)))
                amount = max(1, int(operation.get("amount", 1)))
                sheet.insert_cols(index, amount) if action == "insert_columns" else sheet.delete_cols(index, amount)
            elif action == "merge_cells":
                _sheet(book, str(operation.get("sheet") or book.active.title)).merge_cells(str(operation.get("range", "")).strip())
            elif action == "unmerge_cells":
                _sheet(book, str(operation.get("sheet") or book.active.title)).unmerge_cells(str(operation.get("range", "")).strip())
            elif action == "freeze_panes":
                sheet = _sheet(book, str(operation.get("sheet") or book.active.title))
                value = str(operation.get("cell", "")).strip().upper()
                sheet.freeze_panes = value or None
            elif action == "set_auto_filter":
                sheet = _sheet(book, str(operation.get("sheet") or book.active.title))
                sheet.auto_filter.ref = str(operation.get("range", "")).strip() or None
            elif action == "set_column_width":
                sheet = _sheet(book, str(operation.get("sheet") or book.active.title))
                column = str(operation.get("column", "A")).strip().upper()
                sheet.column_dimensions[column].width = float(operation.get("width", 12))
            elif action == "set_row_height":
                sheet = _sheet(book, str(operation.get("sheet") or book.active.title))
                row = max(1, int(operation.get("row", 1)))
                sheet.row_dimensions[row].height = float(operation.get("height", 15))
            elif action == "add_table":
                sheet = _sheet(book, str(operation.get("sheet") or book.active.title))
                ref = str(operation.get("range", "")).strip().upper()
                name = str(operation.get("name", "Table1")).strip()
                if not ref:
                    raise ValueError("add_table requires range")
                if not _TABLE_NAME.match(name):
                    raise ValueError("Excel table name must start with a letter or underscore and contain only letters, numbers, underscore or period")
                if any(name == table.name for worksheet in book.worksheets for table in worksheet.tables.values()):
                    raise ValueError(f"Excel table already exists: {name}")
                table = Table(displayName=name, ref=ref)
                table.tableStyleInfo = TableStyleInfo(
                    name=str(operation.get("style") or "TableStyleMedium2"),
                    showFirstColumn=False,
                    showLastColumn=False,
                    showRowStripes=True,
                    showColumnStripes=False,
                )
                sheet.add_table(table)
            elif action == "sort":
                sheet = _sheet(book, str(operation.get("sheet") or book.active.title))
                target = str(operation.get("range") or "").upper()
                if not target:
                    raise ValueError("sort requires range")
                min_col, min_row, max_col, max_row = range_boundaries(target)
                column = max(1, int(operation.get("column", 1)))
                if column > (max_col - min_col + 1):
                    raise ValueError("sort column is outside the selected range")
                has_header = bool(operation.get("header", True))
                reverse = str(operation.get("direction") or "ascending") == "descending"
                start_row = min_row + (1 if has_header else 0)
                source_rows = list(sheet.iter_rows(min_row=start_row, max_row=max_row, min_col=min_col, max_col=max_col))
                row_payloads = [
                    [(copy(cell._style), cell.number_format, cell.value) for cell in row]
                    for row in source_rows
                ]
                row_payloads.sort(key=lambda row: (row[column - 1][2] is None, str(row[column - 1][2]).casefold()), reverse=reverse)
                for row_offset, values in enumerate(row_payloads):
                    for column_offset, (style_value, number_format, value) in enumerate(values):
                        cell = sheet.cell(row=start_row + row_offset, column=min_col + column_offset)
                        cell.value = value
                        cell._style = copy(style_value)
                        cell.number_format = number_format
            elif action == "chart":
                sheet = _sheet(book, str(operation.get("sheet") or book.active.title))
                target = str(operation.get("range") or "").upper()
                if not target:
                    raise ValueError("chart requires range")
                min_col, min_row, max_col, max_row = range_boundaries(target)
                kind = str(operation.get("chart_type") or "bar")
                chart = LineChart() if kind == "line" else PieChart() if kind == "pie" else BarChart()
                categories = bool(operation.get("categories", True))
                data = Reference(sheet, min_col=min_col + (1 if categories else 0), max_col=max_col, min_row=min_row, max_row=max_row)
                chart.add_data(data, titles_from_data=True)
                if categories:
                    chart.set_categories(Reference(sheet, min_col=min_col, min_row=min_row + 1, max_row=max_row))
                chart.title = str(operation.get("title") or "Chart")
                chart.style = 10
                anchor = str(operation.get("anchor") or f"{get_column_letter(max_col + 2)}{min_row}")
                sheet.add_chart(chart, anchor)
            else:
                raise ValueError(f"Unsupported XLSX edit action: {action}")
        book.save(path)
    finally:
        book.close()


def compare_xlsx(before: Path, after: Path, limit: int = 500) -> dict[str, Any]:
    left = load_workbook(before, data_only=False, read_only=True)
    right = load_workbook(after, data_only=False, read_only=True)
    changes: list[dict[str, Any]] = []
    try:
        for sheet_name in sorted(set(left.sheetnames) | set(right.sheetnames)):
            if sheet_name not in left.sheetnames:
                changes.append({"sheet": sheet_name, "change": "worksheet added"})
                continue
            if sheet_name not in right.sheetnames:
                changes.append({"sheet": sheet_name, "change": "worksheet removed"})
                continue
            a, b = left[sheet_name], right[sheet_name]
            for row in range(1, max(a.max_row, b.max_row) + 1):
                for column in range(1, max(a.max_column, b.max_column) + 1):
                    old, new = a.cell(row, column).value, b.cell(row, column).value
                    if old != new:
                        changes.append({"sheet": sheet_name, "cell": f"{get_column_letter(column)}{row}", "before": old, "after": new})
                    if len(changes) >= limit:
                        return {"changes": changes, "truncated": True}
        return {"changes": changes, "truncated": False}
    finally:
        left.close(); right.close()


def validate_formulas(path: Path) -> dict[str, Any]:
    book = load_workbook(path, data_only=False, read_only=True)
    issues: list[dict[str, Any]] = []
    formulas = 0
    known = set(book.sheetnames)
    try:
        for sheet in book.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if not isinstance(cell.value, str) or not cell.value.startswith("="):
                        continue
                    formulas += 1
                    try:
                        Tokenizer(cell.value)
                        for quoted, bare in re.findall(r"(?:'([^']+)'|([A-Za-z_][^'!\[\]]*))!", cell.value):
                            referenced = (quoted or bare).strip()
                            if referenced and referenced not in known:
                                issues.append({"sheet": sheet.title, "cell": cell.coordinate, "formula": cell.value, "issue": f"Missing worksheet: {referenced}"})
                        if "#REF!" in cell.value:
                            issues.append({"sheet": sheet.title, "cell": cell.coordinate, "formula": cell.value, "issue": "Broken reference"})
                    except Exception as exc:
                        issues.append({"sheet": sheet.title, "cell": cell.coordinate, "formula": cell.value, "issue": str(exc)})
        return {"formulas": formulas, "issues": issues, "valid": not issues}
    finally:
        book.close()


def range_context(path: Path, sheet_name: str, cell_range: str, max_cells: int = 500) -> str:
    book = load_workbook(path, data_only=False, read_only=True)
    try:
        sheet = _sheet(book, sheet_name)
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
        lines = [f"WORKBOOK RANGE: {path.name} / {sheet.title}!{cell_range}"]
        count = 0
        for row in sheet.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
            values = []
            for cell in row:
                values.append(f"{cell.coordinate}={cell.value!s}")
                count += 1
                if count >= max_cells:
                    break
            lines.append(" | ".join(values))
            if count >= max_cells:
                lines.append("[selection truncated]")
                break
        return "\n".join(lines)
    finally:
        book.close()


def recalculate(path: Path) -> dict[str, Any]:
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        return {"recalculated": False, "engine": "unavailable", "message": "LibreOffice is not installed; formulas remain set to recalculate when opened."}
    with tempfile.TemporaryDirectory(prefix="olladex-recalc-") as folder:
        result = subprocess.run(
            [executable, "--headless", "--convert-to", "xlsx", "--outdir", folder, str(path)],
            capture_output=True,
            text=True,
            timeout=120,
        )
        output = Path(folder) / f"{path.stem}.xlsx"
        converted = result.returncode == 0 and output.exists()
        if converted:
            shutil.copy2(output, path)
    return {"recalculated": converted, "engine": "libreoffice", "message": (result.stdout or result.stderr).strip()}


def export_csv(path: Path, sheet_name: str, output: Path) -> Path:
    book = load_workbook(path, data_only=False, read_only=True)
    try:
        sheet = _sheet(book, sheet_name)
        output.parent.mkdir(parents=True, exist_ok=True)
        with output.open("w", newline="", encoding="utf-8-sig") as handle:
            csv.writer(handle).writerows(sheet.iter_rows(values_only=True))
        return output
    finally:
        book.close()


def import_csv(path: Path, csv_path: Path, sheet_name: str, start_cell: str = "A1") -> dict[str, Any]:
    book = load_workbook(path, data_only=False, keep_vba=path.suffix.lower() == ".xlsm")
    try:
        sheet = _sheet(book, sheet_name)
        min_col, min_row, _, _ = range_boundaries(start_cell)
        rows = 0
        columns = 0
        with csv_path.open(newline="", encoding="utf-8-sig") as handle:
            for row_offset, values in enumerate(csv.reader(handle)):
                rows += 1
                columns = max(columns, len(values))
                for column_offset, value in enumerate(values):
                    sheet.cell(row=min_row + row_offset, column=min_col + column_offset, value=_coerce(value))
        book.save(path)
        return {"action": "import_csv", "sheet": sheet_name, "start_cell": start_cell, "rows": rows, "columns": columns}
    finally:
        book.close()


def _sheet(book, name: str):
    if name not in book.sheetnames:
        raise ValueError(f"Worksheet not found: {name}")
    return book[name]


def _sheet_name(value: Any) -> str:
    name = str(value or "").strip()
    if not name:
        raise ValueError("Worksheet name is required")
    return name[:31]


def _cell(operation: dict[str, Any], key: str = "cell") -> str:
    cell = str(operation.get(key, "")).strip().upper()
    if not cell:
        raise ValueError(f"{key} requires a cell address")
    return cell


def _hex(value: Any) -> str:
    text = str(value or "").strip().lstrip("#").upper()
    if not text:
        return ""
    if len(text) not in {6, 8} or any(character not in "0123456789ABCDEF" for character in text):
        raise ValueError("Colour must be a 6 or 8 digit hexadecimal value")
    return text


def _coerce(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    if value.startswith("="):
        return value
    if value.casefold() in {"true", "false"}:
        return value.casefold() == "true"
    if re.fullmatch(r"-?\d+", value):
        return int(value)
    if re.fullmatch(r"-?(?:\d+\.\d*|\d*\.\d+)", value):
        return float(value)
    return value


def _font_rgb(cell) -> str:
    if cell.font.color and cell.font.color.type == "rgb" and cell.font.color.rgb:
        return str(cell.font.color.rgb)
    return ""
